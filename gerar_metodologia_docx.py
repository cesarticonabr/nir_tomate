"""Gera METODOLOGIA.docx com a explicacao da metodologia"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = Document()
for _ in range(5): doc.add_paragraph()
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('Metodologia'); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = RGBColor(0,51,102)
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run('Predicao de Parametros de Tomate por Espectroscopia NIRS'); r.font.size = Pt(14); r.font.color.rgb = RGBColor(80,80,80)
d = doc.add_paragraph(); d.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = d.add_run('Julho de 2026'); r.font.size = Pt(12); r.font.color.rgb = RGBColor(120,120,120)
doc.add_page_break()

def heading(text, level=1): doc.add_heading(text, level=level)
def para(text):
    doc.add_paragraph(text)
def bullet(text):
    doc.add_paragraph(text, style='List Bullet')
def table(headers, rows):
    t = doc.add_table(rows=len(rows)+1, cols=len(headers), style='Light Shading Accent 1')
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(headers):
        t.cell(0,j).text = h
        for pp in t.cell(0,j).paragraphs:
            for rr in pp.runs: rr.bold = True
    for i,row in enumerate(rows):
        for j,val in enumerate(row):
            t.cell(i+1,j).text = str(val)

heading('1. Dados')
para('Fonte: Planilha Capturas_tomate_NIRS_16jun26_corrigido.xlsx com 2 abas:')
table(['Aba', 'Conteudo', 'Dimensao'], [
    ('Capturas_NIRS', '45 espectros de reflectancia (colunas) x 4200 comprimentos de onda (linhas)', '4200 x 46'),
    ('Dados_analise_destrutiva', '9 parcelas com valores de referencia (AT, Firm, pH, SS, etc.)', '9 x 12'),
])
para('Espectros: 400-2500 nm, passo de 0.5 nm. Cada espectro = 1 fruto.')
para('Estrutura experimental: 3 tratamentos (T1=italiano, T2=cereja) x 3 repeticoes x 5 frutos = 45 frutos.')
para('Dados destrutivos medidos por parcela (1 valor para cada 5 frutos).')

heading('2. Pre-processamento Espectral')
p = doc.add_paragraph()
p.add_run('Standard Normal Variate (SNV): ').bold = True
p.add_run('cada espectro e centrado e escalado. Remove efeitos de linha de base.')
p = doc.add_paragraph()
p.add_run('Savitzky-Golay (1a derivada): ').bold = True
p.add_run('opcional, janela=11, ordem=2. Remove tendencias de linha de base.')

heading('3. Estrutura de Validacao')
para('Problema: 45 espectros mas apenas 9 valores independentes de referencia.')
p = doc.add_paragraph()
p.add_run('Solucao: GroupKFold com grupos = parcelas').bold = True
para('Cada fold: treino com 8 parcelas (40 frutos), teste com 1 parcela (5 frutos).')
para('Os 5 frutos de uma mesma parcela nunca sao separados entre treino e teste.')
para('Validacao conservadora adicional: medias dos 5 frutos por parcela (n=9) + LOOCV.')

heading('4. Pipeline de Machine Learning')
para('Espectros (45x4200) -> SNV -> GroupKFold -> StandardScaler -> PLSRegression -> R2/RMSE/RPD')
p = doc.add_paragraph()
p.add_run('PLS (Partial Least Squares): ').bold = True
p.add_run('padrao ouro para NIR. Cria componentes latentes maximizando covariancia X-y.')
bullet('Lida com alta dimensionalidade (4200 features > 45 amostras)')
bullet('Lida com multicolinearidade entre comprimentos de onda vizinhos')
bullet('Interpretavel (loadings mostram comprimentos importantes)')
p = doc.add_paragraph()
p.add_run('Outros modelos: ').bold = True
p.add_run('Random Forest (lento, mesma performance), SVR (R2 negativo), Ridge/OLS (LOOCV).')

heading('5. Metricas')
table(['Metrica', 'Formula', 'Interpretacao'], [
    ('R2', '1 - S(yi-yi_hat)^2/S(yi-y_bar)^2', 'Proporcao da variancia explicada'),
    ('RMSE', 'sqrt(S(yi-yi_hat)^2/n)', 'Raiz do erro quadratico medio'),
    ('RPD', 'DP(y)/RMSE', 'Relacao de Desempenho Residual'),
])
para('')
p = doc.add_paragraph()
p.add_run('Classificacao RPD (Williams, 2001):').bold = True
table(['RPD', 'Classificacao'], [
    ('> 2.5', 'Excelente'), ('2.0-2.5', 'Bom'), ('1.5-2.0', 'Razoavel'), ('< 1.5', 'Fraco'),
])

heading('6. Resultados')
table(['Grupo', 'Parametro', 'R2', 'RPD', 'Classif'], [
    ('Morfologia', 'Comp', '0.930', '3.79', 'Excelente'),
    ('Morfologia', 'Diam.Equat.', '0.937', '4.00', 'Excelente'),
    ('Morfologia', 'C:D', '0.862', '2.69', 'Excelente'),
    ('Textura', 'Firm', '0.765', '2.06', 'Bom'),
    ('Qualidade', 'SS', '0.751', '2.00', 'Bom'),
    ('Qualidade', 'pH', '0.595', '1.57', 'Razoavel'),
    ('Qualidade', 'Vit. C', '0.236', '1.14', 'Fraco'),
    ('Qualidade', 'AT', '0.093', '1.05', 'Fraco'),
    ('Qualidade', 'SS/AT', '< 0', '0.95', 'Inviavel'),
])

heading('7. Ferramentas')
table(['Ferramenta', 'Funcao'], [
    ('Python 3.10', 'Linguagem principal'),
    ('pandas/numpy', 'Manipulacao de dados'),
    ('scikit-learn', 'Machine Learning (PLS, RF, SVR, PCA, GroupKFold)'),
    ('scipy', 'Pre-processamento e estatistica'),
    ('matplotlib', 'Visualizacao'),
    ('python-docx', 'Relatorios Word'),
    ('Git/GitHub', 'Versionamento'),
])

output = 'METODOLOGIA.docx'
doc.save(output)
print(f'Salvo: {output} ({os.path.getsize(output)/1024:.0f} KB)')
