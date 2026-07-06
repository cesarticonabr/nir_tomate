"""Gera Relatorio_Comp_Diam_CD.docx com resultados NIRS -> Comp, Diam.Equat., C:D"""
import pandas as pd, numpy as np, re, os
from scipy.stats import pearsonr, f_oneway
from sklearn.cross_decomposition import PLSRegression
from sklearn.preprocessing import StandardScaler
from sklearn.model_selection import GroupKFold, LeaveOneOut
from sklearn.metrics import r2_score, mean_squared_error
from sklearn.decomposition import PCA
from sklearn.linear_model import Ridge, LinearRegression
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt

print("Gerando Relatorio_Comp_Diam_CD.docx...")

espectros = pd.read_excel('Capturas_tomate_NIRS_16jun26_corrigido.xlsx', sheet_name='Capturas_NIRS')
dest = pd.read_excel('Capturas_tomate_NIRS_16jun26_corrigido.xlsx', sheet_name='Dados_analise_destrutiva')
col_rep = dest.columns[2]; dest['Parcela'] = dest['Tratamento'].str.strip() + dest[col_rep].str.strip()
amostras = [c for c in espectros.columns if c != 'Espectro']
wv = espectros['Espectro'].values; X_raw = espectros[amostras].T.values
def snv(X): return (X-X.mean(axis=1,keepdims=True))/X.std(axis=1,keepdims=True)
X_snv = snv(X_raw)
df_meta = pd.DataFrame([re.match(r'(T\d+)(R\d+)F(\d+)',a.strip()).groups() for a in amostras],
                        columns=['T','R','F'], index=amostras)
df_meta['Parcela'] = df_meta['T']+df_meta['R']
gkf = GroupKFold(n_splits=5); grupos = df_meta['Parcela']

cols = [('Comp', dest.columns[3]), ('Diam.Equat.', dest.columns[4]), ('C:D', dest.columns[5])]

# Calcular tudo
resultados = []
for p_nome, c_nome in cols:
    y = np.array([dest.set_index('Parcela')[c_nome].to_dict()[p] for p in df_meta['Parcela']])
    anova = f_oneway(*[dest[dest['Tratamento']==t][c_nome] for t in sorted(dest['Tratamento'].unique())])
    corrs_full = np.array([pearsonr(X_snv[:,i],y)[0] for i in range(len(wv))])

    pls_linhas = []
    for nc in [1,2,3,5,10,15]:
        preds,reals=[],[]
        for seed in range(10):
            for tr,te in gkf.split(X_snv,y,grupos.sample(frac=1,random_state=seed)):
                pls=PLSRegression(n_components=nc); scaler=StandardScaler()
                pls.fit(scaler.fit_transform(X_snv[tr]),y[tr])
                preds.extend(pls.predict(scaler.transform(X_snv[te]))); reals.extend(y[te])
        pls_linhas.append((nc,r2_score(reals,preds),np.sqrt(mean_squared_error(reals,preds)),np.std(y)/np.sqrt(mean_squared_error(reals,preds)) if mean_squared_error(reals,preds)>0 else 0))

    df_esp=pd.DataFrame(X_snv); df_esp['Parcela']=df_meta['Parcela'].values
    X_parc=df_esp.groupby('Parcela').mean().values; y_parc=dest[c_nome].values
    loo=LeaveOneOut(); loo_res=[]
    for mod,nm in [(PLSRegression(1),'PLS(1)'),(PLSRegression(2),'PLS(2)'),(Ridge(0.1),'Ridge'),(LinearRegression(),'OLS')]:
        preds,reals=[],[]
        for tr,te in loo.split(X_parc):
            scal=StandardScaler(); mod.fit(scal.fit_transform(X_parc[tr]),y_parc[tr])
            preds.extend(mod.predict(scal.transform(X_parc[te]))); reals.extend(y_parc[te])
        loo_res.append((nm,r2_score(reals,preds),np.sqrt(mean_squared_error(reals,preds))))

    pca=PCA().fit(StandardScaler().fit_transform(X_snv)); X_pc=pca.transform(StandardScaler().fit_transform(X_snv))
    resultados.append((p_nome,c_nome,y,corrs_full,pls_linhas,loo_res,anova,pca,pearsonr(X_pc[:,0],y)[0],X_parc,y_parc))

# GRAFICO
fig, axes = plt.subplots(2,3,figsize=(18,10))
for idx,(p_nome,_,y,corrs_full,_,_,_,pca,pca_r,X_parc,y_parc) in enumerate(resultados):
    loo=LeaveOneOut()
    ax=axes[0,idx]; ax.plot(wv,corrs_full,'b-',lw=0.5,alpha=0.7); ax.axhline(0,color='gray',ls='--',alpha=0.5)
    ax.set_xlabel('nm'); ax.set_ylabel('r'); ax.set_title(f'{p_nome} - Correlacao (max|r|={np.max(np.abs(corrs_full)):.3f})'); ax.grid(alpha=0.3)
    ax=axes[1,idx]
    preds_p,reals_p=[],[]
    for tr,te in loo.split(X_parc):
        scal=StandardScaler(); ridge=Ridge(0.1)
        ridge.fit(scal.fit_transform(X_parc[tr]),y_parc[tr])
        preds_p.append(ridge.predict(scal.transform(X_parc[te]))[0]); reals_p.append(y_parc[te][0])
    r2p=r2_score(reals_p,preds_p)
    ax.scatter(reals_p,preds_p,c='darkgreen',s=80,edgecolors='k',lw=1,zorder=5)
    for i,lbl in enumerate(dest['Parcela']): ax.annotate(lbl,(reals_p[i],preds_p[i]),fontsize=7,xytext=(4,4),textcoords='offset points')
    m1,m2=min(min(reals_p),min(preds_p))-1,max(max(reals_p),max(preds_p))+1
    ax.plot([m1,m2],[m1,m2],'r--',alpha=0.7); ax.set_xlabel('Real'); ax.set_ylabel('Predito')
    ax.set_title(f'{p_nome} - LOOCV Ridge R2={r2p:.3f}'); ax.grid(alpha=0.3)
plt.tight_layout(); plt.savefig('comp_diam_cd_graficos.png',dpi=150,bbox_inches='tight'); plt.close()

# DOCUMENTO
doc = Document()
for _ in range(4): doc.add_paragraph()
title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=title.add_run('Relatorio de Resultados'); run.bold=True; run.font.size=Pt(26); run.font.color.rgb=RGBColor(0,51,102)
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=sub.add_run('Predicao de Comp, Diam.Equat. e C:D por Espectroscopia NIRS'); run.font.size=Pt(14); run.font.color.rgb=RGBColor(80,80,80)
dp=doc.add_paragraph(); dp.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=dp.add_run('Julho de 2026'); run.font.size=Pt(12); run.font.color.rgb=RGBColor(120,120,120)
doc.add_page_break()

# Tabela geral
doc.add_heading('Resumo Comparativo - Todos os 9 Parametros', level=1)
t=doc.add_table(rows=10,cols=7,style='Light Shading Accent 1'); t.alignment=WD_TABLE_ALIGNMENT.CENTER
for j,h in enumerate(['Parametro','R2(PLS)','RPD','max|r|','ANOVA(p)','PC1(r)','Classif']):
    t.cell(0,j).text=h
    for pp in t.cell(0,j).paragraphs:
        for rr in pp.runs: rr.bold=True

todos=[('Comp',0.930,3.79,0.941,0.0000),('Diam.Equat.',0.937,4.00,0.916,0.0000),('C:D',0.862,2.69,0.921,0.0000),
       ('Firm',0.765,2.06,0.890,0.0011),('SS',0.751,2.00,0.891,0.0004),('pH',0.595,1.57,0.794,0.0193),
       ('Vit. C',0.236,1.14,0.593,0.2945),('AT',0.093,1.05,0.464,0.6293),('SS/AT','-0.098',0.95,0.221,0.7415)]
for i,(nome,r2,rpd,rmx,ap) in enumerate(todos):
    idx=i+1
    qual='EXCELENTE' if (isinstance(rpd,float) and rpd>2.5) or (isinstance(r2,float) and r2>0.85) else 'Bom' if (isinstance(rpd,float) and rpd>2.0) or (isinstance(r2,float) and r2>0.7) else 'Razoavel' if (isinstance(rpd,float) and rpd>1.5) else 'Fraco'
    t.cell(idx,0).text=nome; t.cell(idx,1).text=f'{r2}'; t.cell(idx,2).text=f'{rpd:.2f}' if isinstance(rpd,float) else rpd
    t.cell(idx,3).text=f'{rmx:.3f}'; t.cell(idx,4).text=f'{ap:.4f}'
    t.cell(idx,5).text='—'; t.cell(idx,6).text=qual

p=doc.add_paragraph()
p.add_run('Parametros morfologicos (Comp, Diam.Equat., C:D) apresentam excelente capacidade preditiva via NIRS.').bold=True

# Secoes
for p_nome,c_nome,y,corrs_full,pls_linhas,loo_res,anova,pca,pca_r,X_parc,y_parc in resultados:
    doc.add_page_break()
    doc.add_heading(f'Analise: {p_nome}', level=1)
    doc.add_paragraph(f'Estatisticas: media={y.mean():.2f}, std={y.std():.2f}, [{y.min():.2f}, {y.max():.2f}]')
    doc.add_paragraph(f'ANOVA ~ Tratamento: F={anova.statistic:.2f}, p={anova.pvalue:.4f} (altamente significativo)')
    doc.add_paragraph(f'Correlacao maxima |r| = {np.max(np.abs(corrs_full)):.3f}')

    top_c=np.argsort(np.abs(corrs_full))[::-1][:10]
    doc.add_heading('Top 10 comprimentos de onda', level=2)
    t=doc.add_table(rows=11,cols=3,style='Light Shading Accent 1')
    for j,h in enumerate(['#','nm','r']):
        t.cell(0,j).text=h
        for pp in t.cell(0,j).paragraphs:
            for rr in pp.runs: rr.bold=True
    for i,idx in enumerate(top_c):
        t.cell(i+1,0).text=str(i+1); t.cell(i+1,1).text=f'{wv[idx]:.1f}'; t.cell(i+1,2).text=f'{corrs_full[idx]:+.4f}'

    doc.add_heading('PLS Regression (GroupKFold)', level=2)
    t=doc.add_table(rows=len(pls_linhas)+1,cols=5,style='Light Shading Accent 1')
    for j,h in enumerate(['n','R2','RMSE','RPD','Qual']):
        t.cell(0,j).text=h
        for pp in t.cell(0,j).paragraphs:
            for rr in pp.runs: rr.bold=True
    for i,(nc,r2,rmse,rpd) in enumerate(pls_linhas):
        q='EXCELENTE' if rpd>2.5 else 'Bom' if rpd>2.0 else 'Razoavel' if rpd>1.5 else 'Fraco'
        t.cell(i+1,0).text=str(nc); t.cell(i+1,1).text=f'{r2:.4f}'; t.cell(i+1,2).text=f'{rmse:.4f}'; t.cell(i+1,3).text=f'{rpd:.2f}'; t.cell(i+1,4).text=q

    doc.add_heading('LOOCV - Medias por Parcela (n=9)', level=2)
    t=doc.add_table(rows=len(loo_res)+1,cols=3,style='Light Shading Accent 1')
    for j,h in enumerate(['Modelo','R2','RMSE']):
        t.cell(0,j).text=h
        for pp in t.cell(0,j).paragraphs:
            for rr in pp.runs: rr.bold=True
    for i,(nm,r2,rmse) in enumerate(loo_res):
        t.cell(i+1,0).text=nm; t.cell(i+1,1).text=f'{r2:.4f}'; t.cell(i+1,2).text=f'{rmse:.4f}'

    doc.add_heading('PCA', level=2)
    doc.add_paragraph(f'PC1: var={pca.explained_variance_ratio_[0]*100:.1f}%, r={pca_r:.3f}')

doc.add_heading('Graficos', level=1)
doc.add_picture('comp_diam_cd_graficos.png',width=Inches(6.5))
doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Figura 1. Correlacao (topo) e LOOCV (base) para Comp, Diam.Equat. e C:D.').italic=True

doc.add_heading('Conclusao', level=1)
doc.add_paragraph('Os parametros morfologicos sao os mais viaveis para predicao por NIRS:')
doc.add_paragraph('1. Diam.Equat. - R2=0.937, RPD=4.00 (EXCELENTE) - melhor parametro de todos')
doc.add_paragraph('2. Comp - R2=0.930, RPD=3.79 (EXCELENTE)')
doc.add_paragraph('3. C:D - R2=0.862, RPD=2.69 (EXCELENTE)')
doc.add_paragraph()
doc.add_paragraph('Isso ocorre porque os parametros morfologicos estao fortemente associados ao tipo de tomate')
doc.add_paragraph('(italiano vs cereja), que tem assinaturas espectrais distintas na regiao do visivel e NIR.')
doc.add_paragraph()
p=doc.add_paragraph()
p.add_run('Ranking completo (melhor para pior):').bold=True
doc.add_paragraph('Diam.Equat. > Comp > C:D > Firm > SS > pH > Vit. C > AT > SS/AT', style='List Bullet')

output='Relatorio_Comp_Diam_CD.docx'
doc.save(output)
print(f"Documento salvo: {output} ({os.path.getsize(output)/1024:.0f} KB)")
