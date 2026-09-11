"""Gera RESULTADOS_O4.docx - resultados + interpretacao da posicao O4."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = Document()
for _ in range(5):
    doc.add_paragraph()
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('Resultados - Analise NIRS da posicao O4'); r.bold = True
r.font.size = Pt(23); r.font.color.rgb = RGBColor(0, 51, 102)
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run('stem-end to the left - Predicao das 16 caracteristicas do bloco P2')
r.font.size = Pt(12); r.font.color.rgb = RGBColor(80, 80, 80)
d = doc.add_paragraph(); d.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = d.add_run('Setembro de 2026'); r.font.size = Pt(11); r.font.color.rgb = RGBColor(120, 120, 120)
doc.add_page_break()


def heading(x, level=1): doc.add_heading(x, level=level)
def para(x): doc.add_paragraph(x)
def bold_para(pfx, rest=''):
    p = doc.add_paragraph(); p.add_run(pfx).bold = True
    if rest: p.add_run(rest)
def table(headers, rows, fs=8):
    tb = doc.add_table(rows=len(rows) + 1, cols=len(headers), style='Light Shading Accent 1')
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        tb.cell(0, j).text = h
        for pp in tb.cell(0, j).paragraphs:
            for rr in pp.runs: rr.bold = True
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            tb.cell(i + 1, j).text = str(val)
    for row in tb.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                for rr in pp.runs: rr.font.size = Pt(fs)
    doc.add_paragraph('')


heading('0. Nota sobre abreviaturas')
para('Glossario completo na Secao 2 de RESULTADOS_O1.docx. r_within: correlacao obs x pred apos remover a '
     'media de cada cultivar (~0 -> so separa cultivares). R2_loco: R2 ao treinar em 2 cultivares e prever o '
     '3o (muito negativo -> nao generaliza). p_perm: p do teste de permutacao.')

heading('1. Metodo (identico as demais posicoes)')
para('PLSR 1-6 LV | grid de 5 pre-processamentos | 450-2450 nm | Leave-One-Fruit-Out (21 folds) | permutacao '
     '(99x) | leave-one-cultivar-out | correlacao dentro de cultivar. n = 21 frutos (TI/TH/TS, 7 cada).')

heading('2. Tabela mestra - O4 (ordenada por R2cv)')
table(['Caracteristica','Pre-proc.','nLV','R2cv','RMSECV','RPD','RPIQ','r_within','R2_loco','p_perm','Williams'], [
    ('pH','SNV+D2','6','0,77','0,033','2,15','3,92','+0,05','-0,91','0,01','bom'),
    ('Fruit diameter','SNV+D1','5','0,75','4,18','2,04','3,21','+0,50','-0,50','0,01','bom'),
    ('Fruit length','SNV+D2','2','0,73','4,38','1,98','3,86','+0,38','-0,57','0,01','razoavel'),
    ('MF (massa fresca)','SNV+D1','5','0,70','20,8','1,88','3,47','+0,45','-1,06','0,01','razoavel'),
    ('C/D','SNV+D2','3','0,64','0,105','1,71','2,95','+0,27','-1,06','0,01','razoavel'),
    ('a*/b*','SNV','1','0,35','0,115','1,27','1,40','+0,45','+0,03','0,01','fraco'),
    ('Hue','SNV','1','0,26','2,95','1,19','1,15','+0,38','-0,04','0,01','fraco'),
    ('L*','SNV','1','0,26','2,57','1,19','1,72','+0,30','-0,01','0,01','fraco'),
    ('Vit. C','SNV+D2','2','0,11','18,3','1,09','1,16','-0,20','-1,15','0,05','fraco'),
    ('b*','SNV+D1','1','0,02','2,71','1,04','1,16','-0,09','-0,30','0,07','fraco'),
    ('Firmness','SNV+D2','1','-0,04','1,233','1,01','1,62','-0,09','+0,14','0,16','fraco'),
    ('a*','bruto','1','-0,04','4,42','1,01','0,59','+0,08','-0,36','0,06','fraco'),
    ('Chroma','bruto','1','-0,05','4,28','1,00','0,65','+0,04','-0,23','0,08','fraco'),
    ('Titratable acidity','SNV+D1','1','-0,08','0,067','0,98','0,99','-0,05','-0,25','0,16','fraco'),
    ('SS','bruto','1','-0,09','0,414','0,98','1,45','-0,41','-0,26','0,19','fraco'),
    ('SS/AT','bruto','1','-0,15','1,993','0,96','0,79','-0,20','-0,61','0,45','fraco'),
])

heading('3. Interpretacao')
heading('3.1 O4 volta ao padrao de confundimento (nivel O1/O2), sem repetir o achado de O3', level=2)
para('Nenhuma caracteristica em O4 tem R2_loco claramente positivo. As morfologicas (pH, diametro, '
     'comprimento, MF, C/D) tem R2cv alto (0,64-0,77) mas R2_loco sempre negativo (-0,50 a -1,06) - menos '
     'catastrofico que O2, mas sem generalizar. Diameter tem o melhor r_within isolado da serie para essa '
     'variavel (+0,50), porem ainda com R2_loco -0,50: sinal parcial, nao calibracao limpa.')
heading('3.2 O achado de O3 (L*) nao se repete em O4', level=2)
para('L* em O4: R2cv 0,26, r_within +0,30, R2_loco -0,01 (praticamente zero) - muito abaixo do resultado de '
     'O3 (R2cv 0,79 / +0,82 / +0,80). Confirma que o sinal de L* encontrado em O3 e especifico daquela '
     'orientacao de captura, nao uma propriedade geral do fruto.')
heading('3.3 a*/b* e Hue: sinal residual, mas mais fraco que em O3', level=2)
para('a*/b*: r_within +0,45 (proximo do valor de O3, +0,53) mas R2_loco apenas +0,03 (quase zero, contra '
     '+0,32 em O3). Hue: r_within +0,38 mas R2_loco -0,04. Direcao do sinal e a mesma de O3, porem mais fraca '
     'e sem generalizacao clara.')
heading('3.4 pH e Vit. C - conclusoes reforcadas', level=2)
para('pH: r_within ~ 0 (+0,05) e R2_loco -0,91 - quarta posicao seguida confirmando que o R2cv de pH e '
     'discriminacao de cultivar. Vit. C: r_within negativo (-0,20) em O4, pior que em qualquer outra posicao '
     '- reforca que O1 e a unica posicao com sinal real de Vit. C.')

heading('4. Comparacao O1 x O2 x O3 x O4')
table(['Aspecto', 'O1', 'O2', 'O3', 'O4'], [
    ('Melhor R2_loco positivo', 'Vit. C (+0,21)', 'nenhum', 'L* (+0,80)', 'nenhum'),
    ('Alvos com r_within E R2_loco > 0', '1', '0', '5', '1 (marginal)'),
    ('Confundimento C/D (R2_loco)', '-0,27 (melhor)', '-4,68', '-0,82', '-1,06'),
    ('Confundimento comprimento (R2_loco)', '-3,65', '-29,5', '-0,12 (melhor)', '-0,57'),
])
bold_para('Conclusao: ', 'O3 continua sendo a posicao isolada mais promissora. O4 nao traz achado novo - '
          'repete, em versao atenuada, tanto o confundimento morfologico de O2 quanto os tracos fracos de '
          'cor de O3, sem superar nenhum dos dois.')

heading('5. Proximos passos')
para('Ver relatorio consolidado RESULTADOS_CONSOLIDADO.docx, que reune O1-O4 e a media das 4 posicoes em uma '
     'unica tabela caracteristica x posicao.')

for img, cap in [('resultados_O4/pca_O4.png', 'PCA (SNV) dos espectros O4.'),
                 ('resultados_O4/espectros_O4.png', 'Espectros O4: bruto, SNV e SNV+1a derivada.'),
                 ('resultados_O4/obs_pred_O4.png', 'Observado vs. predito (LOO) - 6 melhores caracteristicas por R2cv.')]:
    if os.path.exists(img):
        doc.add_page_break(); heading('Figura', level=2)
        doc.add_picture(img, width=Inches(6.2)); para(cap)

out = 'RESULTADOS_O4.docx'
doc.save(out)
print('Salvo:', out, f'({os.path.getsize(out)/1024:.0f} KB)')
